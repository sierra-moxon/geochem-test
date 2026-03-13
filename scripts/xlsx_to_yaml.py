"""Convert altona2016experiments.xlsx to YAML conforming to geochem-test schema."""

import argparse
from pathlib import Path

import openpyxl
import yaml


def _qty(value, unit):
    """Build a Quantity dict."""
    return {"quantity_value": float(value), "quantity_unit": unit}


def convert_sheet1(ws):
    """Sheet 1: Extraction experiment, 7-26-16 -> RadonObservation records.

    Layout (data in columns B-E):
      Row 1: title
      Row 2: headers (Minutes since start, Sample name, 222Rn pCi/L, uncertainty)
      Row 3+: data
    """
    observations = []
    for row_idx in range(3, ws.max_row + 1):
        minutes = ws.cell(row=row_idx, column=2).value
        sample_name = ws.cell(row=row_idx, column=3).value
        rn_pci = ws.cell(row=row_idx, column=4).value
        unc = ws.cell(row=row_idx, column=5).value

        if minutes is None or sample_name is None:
            continue

        obs = {
            "sample_name": str(sample_name).strip(),
            "measurement_type": "radon_222",
            "time_since_start": _qty(minutes, "UCUM:min"),
        }
        if rn_pci is not None:
            obs["measurement_value"] = _qty(rn_pci, "UCUM:pCi/L")
        else:
            continue  # skip rows with no measurement
        if unc is not None:
            obs["uncertainty"] = _qty(unc, "UCUM:pCi/L")
        observations.append(obs)

    return {
        "id": "altona-extraction-2016",
        "name": "Altona Well Field 222Rn Extraction Experiment",
        "description": "Radon-222 extraction experiment conducted 7/26/2016 at Altona well field",
        "ber_data_source": "altona2016experiments.xlsx",
        "uri": "https://catalog.data.gov/dataset/altona-well-field-222rn-data",
        "entity_type": ["dataset"],
        "experiment_type": "extraction",
        "observations": observations,
    }


def convert_sheet2(ws):
    """Sheet 2: Tracer Experiment 7-28-17 -> mixed observation records.

    Layout (data in columns B-O):
      Rows 1-3: metadata (C_I, C_Cs, Q) in columns F-H
      Row 4: blank
      Row 5: headers for 3 sub-tables:
        B-D: time/222Rn/2sd
        F-H: time/I(ppb)/Cs(ppb)
        J-O: time/Temp/DO/SpCond/pH/ORP
      Row 6+: data
    """
    # Parse header metadata
    metadata = {}
    ci_val = ws.cell(row=1, column=7).value
    if ci_val is not None:
        metadata["initial_iodine_concentration"] = _qty(ci_val, "UCUM:ppb")
    ccs_val = ws.cell(row=2, column=7).value
    if ccs_val is not None:
        metadata["initial_cesium_concentration"] = _qty(ccs_val, "UCUM:ppm")
    q_val = ws.cell(row=3, column=7).value
    if q_val is not None:
        metadata["flow_rate"] = _qty(q_val, "UCUM:L/min")

    observations = []
    data_start = 6

    for row_idx in range(data_start, ws.max_row + 1):
        # Sub-table 1: Radon (cols B=2, C=3, D=4)
        rn_time = ws.cell(row=row_idx, column=2).value
        rn_val = ws.cell(row=row_idx, column=3).value
        rn_unc = ws.cell(row=row_idx, column=4).value
        if rn_time is not None and rn_val is not None:
            obs = {
                "measurement_type": "radon_222",
                "measurement_value": _qty(rn_val, "UCUM:pCi/L"),
                "time_since_start": _qty(rn_time, "UCUM:min"),
            }
            if rn_unc is not None:
                obs["uncertainty"] = _qty(rn_unc, "UCUM:pCi/L")
            observations.append(obs)

        # Sub-table 2: Tracers (cols F=6, G=7, H=8)
        tr_time = ws.cell(row=row_idx, column=6).value
        i_val = ws.cell(row=row_idx, column=7).value
        cs_val = ws.cell(row=row_idx, column=8).value
        if tr_time is not None:
            if i_val is not None:
                observations.append({
                    "measurement_type": "iodine",
                    "measurement_value": _qty(i_val, "UCUM:ppb"),
                    "time_since_start": _qty(tr_time, "UCUM:min"),
                })
            if cs_val is not None:
                observations.append({
                    "measurement_type": "cesium",
                    "measurement_value": _qty(cs_val, "UCUM:ppb"),
                    "time_since_start": _qty(tr_time, "UCUM:min"),
                })

        # Sub-table 3: Water quality (cols J=10, K=11, L=12, M=13, N=14, O=15)
        wq_time = ws.cell(row=row_idx, column=10).value
        if wq_time is not None:
            wq_cols = [
                (11, "temperature", "UCUM:Cel"),
                (12, "dissolved_oxygen", "UCUM:mg/L"),
                (13, "specific_conductance", "UCUM:uS/cm"),
                (14, "ph", "UCUM:pH"),
                (15, "orp", "UCUM:mV"),
            ]
            for col, mtype, unit in wq_cols:
                val = ws.cell(row=row_idx, column=col).value
                if val is not None:
                    observations.append({
                        "measurement_type": mtype,
                        "measurement_value": _qty(val, unit),
                        "time_since_start": _qty(wq_time, "UCUM:min"),
                    })

    dataset = {
        "id": "altona-tracer-2017",
        "name": "Altona Well Field Tracer Experiment",
        "description": "Tracer experiment conducted 7/28/2017 at Altona well field",
        "ber_data_source": "altona2016experiments.xlsx",
        "uri": "https://catalog.data.gov/dataset/altona-well-field-222rn-data",
        "entity_type": ["dataset"],
        "experiment_type": "tracer",
        "observations": observations,
    }
    if metadata:
        dataset["experiment_metadata"] = metadata

    return dataset


def main():
    parser = argparse.ArgumentParser(description="Convert Altona xlsx to geochem-test YAML")
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("altona2016experiments.xlsx"),
        help="Path to the xlsx file",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("data/generated"),
        help="Output directory for generated YAML files",
    )
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.input, data_only=True)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    # Sheet 1: Extraction experiment
    sheet1 = wb.worksheets[0]
    ds1 = convert_sheet1(sheet1)
    out1 = args.output_dir / "GeochemDataset-extraction.yaml"
    with open(out1, "w") as f:
        yaml.dump(ds1, f, default_flow_style=False, sort_keys=False, allow_unicode=True)
    print(f"Wrote {out1} ({len(ds1['observations'])} observations)")

    # Sheet 2: Tracer experiment
    if len(wb.worksheets) > 1:
        sheet2 = wb.worksheets[1]
        ds2 = convert_sheet2(sheet2)
        out2 = args.output_dir / "GeochemDataset-tracer.yaml"
        with open(out2, "w") as f:
            yaml.dump(ds2, f, default_flow_style=False, sort_keys=False, allow_unicode=True)
        print(f"Wrote {out2} ({len(ds2['observations'])} observations)")


if __name__ == "__main__":
    main()
